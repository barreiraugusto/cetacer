"""Exportación de listados a CSV y Excel.

Las dos salidas comparten el mismo armado de filas, así que la lista que se
descarga en Excel es exactamente la que se ve en pantalla.
"""
import csv
import unicodedata
from datetime import date

from django.http import HttpResponse
from django.utils.text import slugify

from inscripciones.models import Inscripcion

# --- Definición de columnas -------------------------------------------------

COLUMNAS_PARTICIPANTES = [
    ("#", lambda i, n: n),
    ("Apellido", lambda i, n: i.participante.apellido),
    ("Nombre", lambda i, n: i.participante.nombre),
    ("DNI", lambda i, n: i.participante.dni),
    ("Teléfono", lambda i, n: i.participante.telefono),
    ("Correo", lambda i, n: i.participante.email),
    ("Localidad", lambda i, n: i.participante.localidad),
    ("Empresa", lambda i, n: i.participante.empresa.razon_social if i.participante.empresa else ""),
    ("Cat. licencia", lambda i, n: i.participante.categoria_licencia),
    ("Estado", lambda i, n: i.get_estado_display()),
    ("Pago", lambda i, n: i.get_pago_display()),
    ("Monto", lambda i, n: i.monto if i.monto is not None else ""),
    ("Comprobante", lambda i, n: i.comprobante),
    ("Psicofísico", lambda i, n: "Sí" if i.psicofisico_vigente else "No"),
    ("Documentación", lambda i, n: "Sí" if i.documentacion_completa else "No"),
    ("Origen", lambda i, n: i.get_origen_display()),
    ("Inscripto el", lambda i, n: i.creado.strftime("%d/%m/%Y")),
    ("Notas", lambda i, n: i.notas),
]

COLUMNAS_INSCRIPCIONES = [
    ("Curso", lambda i, n: i.comision.curso.nombre),
    ("Categoría", lambda i, n: i.comision.curso.categoria.nombre),
    ("Fecha", lambda i, n: i.comision.fecha_inicio.strftime("%d/%m/%Y")),
    ("Horario", lambda i, n: i.comision.horario_texto),
] + COLUMNAS_PARTICIPANTES[1:]

#: Planilla para imprimir y firmar el día del curso.
COLUMNAS_PLANILLA = [
    ("#", lambda i, n: n),
    ("Apellido y nombre", lambda i, n: i.participante.nombre_completo),
    ("DNI", lambda i, n: i.participante.dni),
    ("Empresa", lambda i, n: i.participante.empresa.razon_social if i.participante.empresa else ""),
    ("Firma", lambda i, n: ""),
    ("Aclaración", lambda i, n: ""),
]


def _filas(inscripciones, columnas):
    for numero, inscripcion in enumerate(inscripciones, start=1):
        yield [obtener(inscripcion, numero) for _, obtener in columnas]


def _nombre_archivo(base):
    limpio = slugify(unicodedata.normalize("NFKD", base))[:80] or "listado"
    return f"{limpio}-{date.today():%Y-%m-%d}"


# --- Salidas ---------------------------------------------------------------

def a_csv(inscripciones, columnas, base_nombre, encabezado_extra=None):
    """CSV con BOM y punto y coma: se abre bien en Excel en español.

    El charset declarado es utf-8 a propósito: con utf-8-sig Django antepone el
    BOM en cada escritura, no sólo al principio. El BOM se escribe una vez, a mano.
    """
    respuesta = HttpResponse(content_type="text/csv; charset=utf-8")
    respuesta["Content-Disposition"] = (
        f'attachment; filename="{_nombre_archivo(base_nombre)}.csv"'
    )
    respuesta.write("﻿")
    escritor = csv.writer(respuesta, delimiter=";")
    for linea in encabezado_extra or []:
        escritor.writerow([linea])
    if encabezado_extra:
        escritor.writerow([])
    escritor.writerow([titulo for titulo, _ in columnas])
    for fila in _filas(inscripciones, columnas):
        escritor.writerow(fila)
    return respuesta


def a_excel(inscripciones, columnas, base_nombre, titulo="Listado", encabezado_extra=None):
    """Planilla .xlsx con encabezado institucional y columnas autoajustadas."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    libro = Workbook()
    hoja = libro.active
    hoja.title = titulo[:31]

    azul = "0F2744"
    azul_claro = "E8EFFF"
    fila_actual = 1

    hoja.cell(row=1, column=1, value="CETACER — Cámara Empresaria del Transporte Automotor de Cargas de Entre Ríos")
    hoja.cell(row=1, column=1).font = Font(bold=True, size=13, color=azul)
    fila_actual = 2
    for linea in encabezado_extra or []:
        hoja.cell(row=fila_actual, column=1, value=linea).font = Font(size=10, color="475467")
        fila_actual += 1
    hoja.cell(row=fila_actual, column=1, value=f"Generado el {date.today():%d/%m/%Y}").font = Font(
        size=9, italic=True, color="667085"
    )
    fila_actual += 2

    fila_encabezado = fila_actual
    borde = Side(style="thin", color="D0D5DD")
    for columna, (titulo_col, _) in enumerate(columnas, start=1):
        celda = hoja.cell(row=fila_encabezado, column=columna, value=titulo_col)
        celda.font = Font(bold=True, color="FFFFFF", size=10)
        celda.fill = PatternFill("solid", fgColor=azul)
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        celda.border = Border(bottom=borde)
    hoja.row_dimensions[fila_encabezado].height = 26

    anchos = [len(t) + 2 for t, _ in columnas]
    fila_actual = fila_encabezado + 1
    for indice, fila in enumerate(_filas(inscripciones, columnas)):
        for columna, valor in enumerate(fila, start=1):
            celda = hoja.cell(row=fila_actual, column=columna, value=valor)
            celda.border = Border(bottom=borde)
            celda.font = Font(size=10)
            if indice % 2 == 1:
                celda.fill = PatternFill("solid", fgColor="F7F9FC")
            largo = len(str(valor)) if valor is not None else 0
            if largo + 2 > anchos[columna - 1]:
                anchos[columna - 1] = min(largo + 2, 45)
        fila_actual += 1

    for columna, ancho in enumerate(anchos, start=1):
        hoja.column_dimensions[get_column_letter(columna)].width = max(ancho, 9)
    hoja.freeze_panes = hoja.cell(row=fila_encabezado + 1, column=1)
    if fila_actual > fila_encabezado + 1:
        hoja.auto_filter.ref = (
            f"A{fila_encabezado}:"
            f"{get_column_letter(len(columnas))}{fila_actual - 1}"
        )

    respuesta = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    respuesta["Content-Disposition"] = (
        f'attachment; filename="{_nombre_archivo(base_nombre)}.xlsx"'
    )
    libro.save(respuesta)
    return respuesta


def exportar(formato, inscripciones, columnas, base_nombre, titulo="Listado", encabezado_extra=None):
    """Punto de entrada único: `formato` es "csv" o "xlsx"."""
    if formato == "csv":
        return a_csv(inscripciones, columnas, base_nombre, encabezado_extra)
    return a_excel(inscripciones, columnas, base_nombre, titulo, encabezado_extra)


def encabezado_comision(comision):
    """Líneas de contexto que encabezan la lista de una comisión."""
    lineas = [
        f"Curso: {comision.curso.nombre}",
        f"Categoría: {comision.curso.categoria.nombre}",
        f"Fecha: {comision.fecha_texto} — {comision.horario_texto}",
        f"Lugar: {comision.lugar_texto}",
        f"Cupo: {comision.cantidad_inscriptos} de {comision.cupo}",
    ]
    if comision.instructor:
        lineas.append(f"Instructor: {comision.instructor.nombre_completo}")
    return lineas


def inscripciones_de(comision, solo_activas=True):
    consulta = comision.inscripciones.select_related(
        "participante", "participante__empresa", "comision", "comision__curso"
    )
    if solo_activas:
        consulta = consulta.exclude(estado=Inscripcion.Estado.CANCELADA)
    return consulta.order_by("participante__apellido", "participante__nombre")
