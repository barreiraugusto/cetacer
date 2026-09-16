"""Vistas del panel de administración de CETACER."""
from datetime import date, timedelta

from django.contrib import messages
from django.db import transaction
from django.db.models import Count, Q, Sum
from django.shortcuts import get_object_or_404, redirect, render
from django.utils.http import url_has_allowed_host_and_scheme

from contenido.models import Pagina
from cuentas.models import DESCRIPCION_ROLES, JERARQUIA, Rol, Usuario
from cursos.models import Categoria, Comision, ConfiguracionSitio, Curso
from inscripciones.models import Empresa, Inscripcion, Participante

from . import exportaciones as exp
from .accesos import (
    GESTION_CATALOGO,
    GESTION_COMISIONES,
    GESTION_INSCRIPCIONES,
    LECTURA_INSTITUCIONAL,
    SOLO_DIRECCION,
    comisiones_visibles,
    requiere_panel,
)
from .forms import (
    CategoriaForm,
    ComisionForm,
    ConfiguracionForm,
    CursoForm,
    EmpresaForm,
    EnlaceFormSet,
    InscribirForm,
    InscripcionForm,
    PaginaForm,
    ParticipanteForm,
    UsuarioForm,
)


def _volver(request, por_defecto):
    """Vuelve a `next` si es una URL propia; si no, al destino por defecto."""
    destino = request.POST.get("next") or request.GET.get("next")
    if destino and url_has_allowed_host_and_scheme(
        destino, allowed_hosts={request.get_host()}, require_https=request.is_secure()
    ):
        return redirect(destino)
    return redirect(por_defecto)


# --- Tablero ---------------------------------------------------------------

@requiere_panel
def inicio(request):
    """Tablero de entrada: lo que hay que mirar hoy, en un vistazo."""
    hoy = date.today()
    comisiones = comisiones_visibles(
        request.user, Comision.objects.select_related("curso", "curso__categoria", "instructor")
    )
    proximas = list(
        comisiones.filter(
            fecha_inicio__gte=hoy,
            estado__in=[Comision.Estado.PUBLICADA, Comision.Estado.EN_CURSO, Comision.Estado.BORRADOR],
        ).order_by("fecha_inicio")[:6]
    )

    inscripciones = Inscripcion.objects.select_related(
        "participante", "comision", "comision__curso"
    )
    if request.user.tiene_rol(Rol.INSTRUCTOR):
        inscripciones = inscripciones.filter(comision__instructor=request.user)

    en_30_dias = hoy + timedelta(days=30)
    activas = inscripciones.exclude(estado=Inscripcion.Estado.CANCELADA)

    indicadores = {
        "comisiones_abiertas": comisiones.filter(
            fecha_inicio__gte=hoy, estado=Comision.Estado.PUBLICADA
        ).count(),
        "inscriptos_proximos": activas.filter(
            comision__fecha_inicio__gte=hoy, comision__fecha_inicio__lte=en_30_dias
        ).count(),
        "pagos_pendientes": activas.filter(
            pago=Inscripcion.Pago.PENDIENTE, comision__fecha_inicio__gte=hoy
        ).count(),
        "solicitudes_web": activas.filter(
            origen="web", estado=Inscripcion.Estado.PREINSCRIPTO
        ).count(),
    }

    # Cosas que requieren una acción concreta, con el enlace a resolverlas.
    avisos = []
    for comision in proximas:
        if comision.estado == Comision.Estado.BORRADOR and comision.dias_faltantes <= 14:
            avisos.append({
                "tipo": "atencion",
                "texto": f"«{comision.curso.nombre}» del {comision.fecha_texto} sigue en borrador y no se ve en la web.",
                "url": f"/panel/comisiones/{comision.pk}/editar/",
                "accion": "Publicar",
            })
        if comision.completa and comision.visible_en_web:
            avisos.append({
                "tipo": "info",
                "texto": f"«{comision.curso.nombre}» del {comision.fecha_texto} llegó al cupo ({comision.cupo}).",
                "url": f"/panel/comisiones/{comision.pk}/",
                "accion": "Ver lista",
            })
    if indicadores["solicitudes_web"]:
        avisos.insert(0, {
            "tipo": "atencion",
            "texto": f"Hay {indicadores['solicitudes_web']} solicitud(es) del formulario web sin confirmar.",
            "url": "/panel/inscripciones/?estado=preinscripta&origen=web",
            "accion": "Revisar",
        })

    # Ocupación por categoría de los próximos 60 días.
    ocupacion = (
        Categoria.objects.filter(
            cursos__comisiones__fecha_inicio__gte=hoy,
            cursos__comisiones__fecha_inicio__lte=hoy + timedelta(days=60),
        )
        .annotate(
            inscriptos=Count(
                "cursos__comisiones__inscripciones",
                filter=~Q(cursos__comisiones__inscripciones__estado=Inscripcion.Estado.CANCELADA),
            ),
            comisiones_total=Count("cursos__comisiones", distinct=True),
        )
        .order_by("-inscriptos")
    )
    tope = max([c.inscriptos for c in ocupacion] or [1]) or 1
    for categoria in ocupacion:
        categoria.barra = round(categoria.inscriptos * 100 / tope)

    return render(request, "panel/inicio.html", {
        "seccion": "inicio",
        "indicadores": indicadores,
        "proximas": proximas,
        "avisos": avisos[:5],
        "ocupacion": ocupacion,
        "ultimas": activas.order_by("-creado")[:8],
    })


@requiere_panel
def agenda(request):
    """Calendario mensual de comisiones."""
    hoy = date.today()
    try:
        anio = int(request.GET.get("anio", hoy.year))
        mes = int(request.GET.get("mes", hoy.month))
        primero = date(anio, mes, 1)
    except (ValueError, TypeError):
        primero = date(hoy.year, hoy.month, 1)

    siguiente = date(primero.year + (primero.month == 12), (primero.month % 12) + 1, 1)
    anterior_mes = primero.month - 1 or 12
    anterior = date(primero.year - (primero.month == 1), anterior_mes, 1)

    comisiones = comisiones_visibles(
        request.user, Comision.objects.select_related("curso", "curso__categoria")
    ).filter(fecha_inicio__gte=primero, fecha_inicio__lt=siguiente).order_by("fecha_inicio")

    por_dia = {}
    for comision in comisiones:
        por_dia.setdefault(comision.fecha_inicio.day, []).append(comision)

    # Grilla de 6 semanas arrancando en lunes.
    desplazamiento = primero.weekday()
    celdas = []
    for indice in range(42):
        dia = primero + timedelta(days=indice - desplazamiento)
        celdas.append({
            "fecha": dia,
            "del_mes": dia.month == primero.month,
            "es_hoy": dia == hoy,
            "comisiones": por_dia.get(dia.day, []) if dia.month == primero.month else [],
        })

    return render(request, "panel/agenda.html", {
        "seccion": "agenda",
        "mes_actual": primero,
        "mes_anterior": anterior,
        "mes_siguiente": siguiente,
        "celdas": celdas,
        "total": comisiones.count(),
    })


# --- Catálogo --------------------------------------------------------------

@requiere_panel
def cursos_lista(request):
    consulta = Curso.objects.select_related("categoria").annotate(
        total_comisiones=Count("comisiones", distinct=True)
    )
    busqueda = request.GET.get("q", "").strip()
    categoria = request.GET.get("categoria", "")
    if busqueda:
        consulta = consulta.filter(Q(nombre__icontains=busqueda) | Q(requisitos__icontains=busqueda))
    if categoria:
        consulta = consulta.filter(categoria_id=categoria)
    return render(request, "panel/cursos.html", {
        "seccion": "cursos",
        "cursos": consulta,
        "categorias": Categoria.objects.all(),
        "q": busqueda,
        "categoria_sel": categoria,
    })


@GESTION_CATALOGO
def curso_editar(request, pk=None):
    curso = get_object_or_404(Curso, pk=pk) if pk else None
    if request.method == "POST":
        formulario = CursoForm(request.POST, instance=curso)
        if formulario.is_valid():
            guardado = formulario.save()
            messages.success(
                request,
                f"Curso «{guardado.nombre}» {'actualizado' if pk else 'creado'}.",
            )
            return redirect("panel:cursos")
        messages.error(request, "Revisá los campos marcados en rojo.")
    else:
        formulario = CursoForm(instance=curso)
    return render(request, "panel/curso_form.html", {
        "seccion": "cursos",
        "form": formulario,
        "curso": curso,
    })


@GESTION_CATALOGO
def curso_eliminar(request, pk):
    curso = get_object_or_404(Curso, pk=pk)
    if request.method != "POST":
        return redirect("panel:curso_editar", pk=pk)
    if curso.comisiones.exists():
        messages.error(
            request,
            "Ese curso tiene comisiones cargadas. Desactivalo en lugar de borrarlo "
            "para no perder el historial de participantes.",
        )
        return redirect("panel:curso_editar", pk=pk)
    nombre = curso.nombre
    curso.delete()
    messages.success(request, f"Curso «{nombre}» eliminado.")
    return redirect("panel:cursos")


@requiere_panel
def categorias_lista(request):
    return render(request, "panel/categorias.html", {
        "seccion": "cursos",
        "categorias": Categoria.objects.annotate(total=Count("cursos")),
    })


@GESTION_CATALOGO
def categoria_editar(request, pk=None):
    categoria = get_object_or_404(Categoria, pk=pk) if pk else None
    if request.method == "POST":
        formulario = CategoriaForm(request.POST, instance=categoria)
        if formulario.is_valid():
            formulario.save()
            messages.success(request, "Categoría guardada.")
            return redirect("panel:categorias")
    else:
        formulario = CategoriaForm(instance=categoria)
    return render(request, "panel/categoria_form.html", {
        "seccion": "cursos",
        "form": formulario,
        "categoria": categoria,
    })


@LECTURA_INSTITUCIONAL
def paginas_lista(request):
    return render(request, "panel/paginas.html", {
        "seccion": "paginas",
        "paginas_cargadas": Pagina.objects.prefetch_related("enlaces"),
    })


@GESTION_CATALOGO
def pagina_editar(request, pk=None):
    """La página y sus enlaces se guardan juntos o no se guarda nada."""
    pagina = get_object_or_404(Pagina, pk=pk) if pk else None
    if request.method == "POST":
        formulario = PaginaForm(request.POST, instance=pagina)
        enlaces = EnlaceFormSet(request.POST, request.FILES, instance=pagina)
        if formulario.is_valid() and enlaces.is_valid():
            with transaction.atomic():
                guardada = formulario.save()
                enlaces.instance = guardada
                enlaces.save()
            messages.success(request, f"Página «{guardada.titulo}» guardada.")
            return redirect("panel:paginas")
        messages.error(request, "Revisá los campos marcados en rojo.")
    else:
        formulario = PaginaForm(instance=pagina)
        enlaces = EnlaceFormSet(instance=pagina)
    return render(request, "panel/pagina_form.html", {
        "seccion": "paginas",
        "form": formulario,
        "enlaces": enlaces,
        "pagina": pagina,
    })


@GESTION_CATALOGO
def pagina_eliminar(request, pk):
    pagina = get_object_or_404(Pagina, pk=pk)
    if request.method != "POST":
        return redirect("panel:pagina_editar", pk=pk)
    titulo = pagina.titulo
    pagina.delete()
    messages.success(request, f"Página «{titulo}» eliminada.")
    return redirect("panel:paginas")


# --- Comisiones ------------------------------------------------------------

@requiere_panel
def comisiones_lista(request):
    consulta = comisiones_visibles(
        request.user,
        Comision.objects.select_related("curso", "curso__categoria", "instructor"),
    ).annotate(
        inscriptos_total=Count(
            "inscripciones", filter=~Q(inscripciones__estado=Inscripcion.Estado.CANCELADA)
        )
    )
    estado = request.GET.get("estado", "")
    cuando = request.GET.get("cuando", "proximas")
    busqueda = request.GET.get("q", "").strip()
    if estado:
        consulta = consulta.filter(estado=estado)
    if cuando == "proximas":
        consulta = consulta.filter(fecha_inicio__gte=date.today())
    elif cuando == "pasadas":
        consulta = consulta.filter(fecha_inicio__lt=date.today()).order_by("-fecha_inicio")
    if busqueda:
        consulta = consulta.filter(curso__nombre__icontains=busqueda)
    return render(request, "panel/comisiones.html", {
        "seccion": "comisiones",
        "comisiones": consulta,
        "estados": Comision.Estado.choices,
        "estado_sel": estado,
        "cuando": cuando,
        "q": busqueda,
    })


@requiere_panel
def comision_detalle(request, pk):
    comision = get_object_or_404(
        comisiones_visibles(
            request.user, Comision.objects.select_related("curso", "curso__categoria", "instructor")
        ),
        pk=pk,
    )
    inscripciones = comision.inscripciones.select_related(
        "participante", "participante__empresa"
    ).order_by("participante__apellido", "participante__nombre")
    resumen = {
        "total": inscripciones.exclude(estado=Inscripcion.Estado.CANCELADA).count(),
        "confirmadas": inscripciones.filter(estado=Inscripcion.Estado.CONFIRMADA).count(),
        "preinscriptas": inscripciones.filter(estado=Inscripcion.Estado.PREINSCRIPTO).count(),
        "pagadas": inscripciones.filter(pago=Inscripcion.Pago.PAGADO).count(),
        "recaudado": inscripciones.filter(pago=Inscripcion.Pago.PAGADO).aggregate(
            total=Sum("monto")
        )["total"] or 0,
    }
    return render(request, "panel/comision_detalle.html", {
        "seccion": "comisiones",
        "comision": comision,
        "inscripciones": inscripciones,
        "resumen": resumen,
        "form_inscribir": InscribirForm(),
        "puede_inscribir": request.user.is_superuser or request.user.alcanza(Rol.RECEPCION),
        "estados_comision": Comision.Estado.choices,
    })


@GESTION_COMISIONES
def comision_editar(request, pk=None):
    comision = get_object_or_404(Comision, pk=pk) if pk else None
    if request.method == "POST":
        formulario = ComisionForm(request.POST, instance=comision)
        if formulario.is_valid():
            guardada = formulario.save()
            messages.success(request, "Comisión guardada.")
            return redirect("panel:comision_detalle", pk=guardada.pk)
        messages.error(request, "Revisá los campos marcados en rojo.")
    else:
        inicial = {}
        curso_id = request.GET.get("curso")
        if curso_id:
            curso = Curso.objects.filter(pk=curso_id).first()
            if curso:
                inicial = {"curso": curso, "cupo": curso.cupo_sugerido}
        formulario = ComisionForm(instance=comision, initial=inicial)
    return render(request, "panel/comision_form.html", {
        "seccion": "comisiones",
        "form": formulario,
        "comision": comision,
    })


@GESTION_COMISIONES
def comision_cambiar_estado(request, pk):
    comision = get_object_or_404(Comision, pk=pk)
    if request.method != "POST":
        return redirect("panel:comision_detalle", pk=pk)
    nuevo = request.POST.get("estado")
    validos = {valor for valor, _ in Comision.Estado.choices}
    if nuevo not in validos:
        messages.error(request, "Ese estado no existe.")
    else:
        comision.estado = nuevo
        comision.save(update_fields=["estado", "actualizado"])
        messages.success(request, f"La comisión pasó a «{comision.get_estado_display()}».")
    return redirect("panel:comision_detalle", pk=pk)


@GESTION_INSCRIPCIONES
def inscribir(request, pk):
    """Alta rápida: si el DNI ya existe reutiliza la persona, si no la crea."""
    comision = get_object_or_404(Comision, pk=pk)
    if request.method != "POST":
        return redirect("panel:comision_detalle", pk=pk)

    formulario = InscribirForm(request.POST)
    if not formulario.is_valid():
        for errores in formulario.errors.values():
            for error in errores:
                messages.error(request, error)
        return redirect("panel:comision_detalle", pk=pk)

    datos = formulario.cleaned_data
    if comision.completa:
        messages.error(
            request,
            f"La comisión ya tiene {comision.cupo} inscriptos. Ampliá el cupo antes de agregar a alguien más.",
        )
        return redirect("panel:comision_detalle", pk=pk)

    with transaction.atomic():
        participante, creado = Participante.objects.get_or_create(
            dni=datos["dni"],
            defaults={
                "apellido": datos["apellido"],
                "nombre": datos["nombre"],
                "telefono": datos["telefono"],
                "email": datos["email"],
                "empresa": datos["empresa"],
            },
        )
        if not creado:
            # Completa los datos que estaban vacíos sin pisar lo ya cargado.
            cambios = []
            for campo in ("telefono", "email"):
                if datos[campo] and not getattr(participante, campo):
                    setattr(participante, campo, datos[campo])
                    cambios.append(campo)
            if datos["empresa"] and not participante.empresa:
                participante.empresa = datos["empresa"]
                cambios.append("empresa")
            if cambios:
                participante.save(update_fields=cambios)

        _, nueva = Inscripcion.objects.get_or_create(
            comision=comision,
            participante=participante,
            defaults={
                "estado": datos["estado"],
                "pago": datos["pago"],
                "origen": "panel",
                "creado_por": request.user,
            },
        )

    if nueva:
        messages.success(
            request,
            f"{participante.nombre_completo} quedó inscripto. "
            f"Quedan {comision.lugares_disponibles} lugares.",
        )
    else:
        messages.info(request, f"{participante.nombre_completo} ya estaba en esta comisión.")
    return redirect("panel:comision_detalle", pk=pk)


@requiere_panel
def tomar_asistencia(request, pk):
    """Marca asistió/ausente para toda la comisión de una sola vez."""
    comision = get_object_or_404(comisiones_visibles(request.user, Comision.objects.all()), pk=pk)
    inscripciones = comision.inscripciones.select_related("participante").exclude(
        estado=Inscripcion.Estado.CANCELADA
    ).order_by("participante__apellido")

    if request.method == "POST":
        presentes = set(request.POST.getlist("presente"))
        actualizadas = []
        for inscripcion in inscripciones:
            inscripcion.estado = (
                Inscripcion.Estado.ASISTIO
                if str(inscripcion.pk) in presentes
                else Inscripcion.Estado.AUSENTE
            )
            actualizadas.append(inscripcion)
        Inscripcion.objects.bulk_update(actualizadas, ["estado"])
        if comision.estado != Comision.Estado.FINALIZADA:
            comision.estado = Comision.Estado.FINALIZADA
            comision.save(update_fields=["estado", "actualizado"])
        messages.success(
            request,
            f"Asistencia registrada: {len(presentes)} presentes de {len(actualizadas)}.",
        )
        return redirect("panel:comision_detalle", pk=pk)

    return render(request, "panel/asistencia.html", {
        "seccion": "comisiones",
        "comision": comision,
        "inscripciones": inscripciones,
    })


# --- Inscripciones, participantes y empresas -------------------------------

def _filtrar_inscripciones(request):
    consulta = Inscripcion.objects.select_related(
        "participante", "participante__empresa", "comision", "comision__curso",
        "comision__curso__categoria",
    )
    if request.user.tiene_rol(Rol.INSTRUCTOR):
        consulta = consulta.filter(comision__instructor=request.user)

    filtros = {
        "q": request.GET.get("q", "").strip(),
        "estado": request.GET.get("estado", ""),
        "pago": request.GET.get("pago", ""),
        "origen": request.GET.get("origen", ""),
        "curso": request.GET.get("curso", ""),
        "desde": request.GET.get("desde", ""),
        "hasta": request.GET.get("hasta", ""),
    }
    if filtros["q"]:
        consulta = consulta.filter(
            Q(participante__apellido__icontains=filtros["q"])
            | Q(participante__nombre__icontains=filtros["q"])
            | Q(participante__dni__icontains=filtros["q"])
        )
    if filtros["estado"]:
        consulta = consulta.filter(estado=filtros["estado"])
    if filtros["pago"]:
        consulta = consulta.filter(pago=filtros["pago"])
    if filtros["origen"]:
        consulta = consulta.filter(origen=filtros["origen"])
    if filtros["curso"]:
        consulta = consulta.filter(comision__curso_id=filtros["curso"])
    if filtros["desde"]:
        consulta = consulta.filter(comision__fecha_inicio__gte=filtros["desde"])
    if filtros["hasta"]:
        consulta = consulta.filter(comision__fecha_inicio__lte=filtros["hasta"])
    return consulta.order_by("-comision__fecha_inicio", "participante__apellido"), filtros


@requiere_panel
def inscripciones_lista(request):
    consulta, filtros = _filtrar_inscripciones(request)
    return render(request, "panel/inscripciones.html", {
        "seccion": "inscripciones",
        "inscripciones": consulta[:300],
        "total": consulta.count(),
        "filtros": filtros,
        "estados": Inscripcion.Estado.choices,
        "pagos": Inscripcion.Pago.choices,
        "cursos": Curso.objects.all(),
        "querystring": request.GET.urlencode(),
    })


@GESTION_INSCRIPCIONES
def inscripcion_editar(request, pk):
    inscripcion = get_object_or_404(
        Inscripcion.objects.select_related("participante", "comision"), pk=pk
    )
    if request.method == "POST":
        formulario = InscripcionForm(request.POST, instance=inscripcion)
        if formulario.is_valid():
            formulario.save()
            messages.success(request, "Inscripción actualizada.")
            return _volver(request, "panel:inscripciones")
    else:
        formulario = InscripcionForm(instance=inscripcion)
    return render(request, "panel/inscripcion_form.html", {
        "seccion": "inscripciones",
        "form": formulario,
        "inscripcion": inscripcion,
    })


@requiere_panel
def participantes_lista(request):
    consulta = Participante.objects.select_related("empresa").annotate(
        total_cursos=Count("inscripciones", distinct=True)
    )
    busqueda = request.GET.get("q", "").strip()
    empresa = request.GET.get("empresa", "")
    if busqueda:
        consulta = consulta.filter(
            Q(apellido__icontains=busqueda)
            | Q(nombre__icontains=busqueda)
            | Q(dni__icontains=busqueda)
        )
    if empresa:
        consulta = consulta.filter(empresa_id=empresa)
    return render(request, "panel/participantes.html", {
        "seccion": "participantes",
        "participantes": consulta[:300],
        "total": consulta.count(),
        "empresas": Empresa.objects.all(),
        "q": busqueda,
        "empresa_sel": empresa,
        "querystring": request.GET.urlencode(),
    })


@GESTION_INSCRIPCIONES
def participante_editar(request, pk=None):
    participante = get_object_or_404(Participante, pk=pk) if pk else None
    if request.method == "POST":
        formulario = ParticipanteForm(request.POST, instance=participante)
        if formulario.is_valid():
            formulario.save()
            messages.success(request, "Participante guardado.")
            return redirect("panel:participantes")
    else:
        formulario = ParticipanteForm(instance=participante)
    historial = participante.inscripciones.select_related(
        "comision", "comision__curso"
    ).order_by("-comision__fecha_inicio") if participante else []
    return render(request, "panel/participante_form.html", {
        "seccion": "participantes",
        "form": formulario,
        "participante": participante,
        "historial": historial,
    })


@requiere_panel
def empresas_lista(request):
    return render(request, "panel/empresas.html", {
        "seccion": "participantes",
        "empresas": Empresa.objects.annotate(total=Count("participantes")),
    })


@GESTION_INSCRIPCIONES
def empresa_editar(request, pk=None):
    empresa = get_object_or_404(Empresa, pk=pk) if pk else None
    if request.method == "POST":
        formulario = EmpresaForm(request.POST, instance=empresa)
        if formulario.is_valid():
            formulario.save()
            messages.success(request, "Empresa guardada.")
            return redirect("panel:empresas")
    else:
        formulario = EmpresaForm(instance=empresa)
    return render(request, "panel/empresa_form.html", {
        "seccion": "participantes",
        "form": formulario,
        "empresa": empresa,
    })


# --- Exportaciones ---------------------------------------------------------

@requiere_panel
def exportar_comision(request, pk):
    """Lista de participantes de una comisión, en Excel o CSV."""
    comision = get_object_or_404(comisiones_visibles(request.user, Comision.objects.all()), pk=pk)
    formato = request.GET.get("formato", "xlsx")
    incluir_canceladas = request.GET.get("canceladas") == "1"
    inscripciones = exp.inscripciones_de(comision, solo_activas=not incluir_canceladas)
    return exp.exportar(
        formato,
        inscripciones,
        exp.COLUMNAS_PARTICIPANTES,
        f"participantes-{comision.curso.nombre}-{comision.fecha_inicio}",
        titulo="Participantes",
        encabezado_extra=exp.encabezado_comision(comision),
    )


@requiere_panel
def planilla_asistencia(request, pk):
    """Planilla para imprimir y hacer firmar el día del curso."""
    comision = get_object_or_404(comisiones_visibles(request.user, Comision.objects.all()), pk=pk)
    inscripciones = exp.inscripciones_de(comision)
    return exp.exportar(
        request.GET.get("formato", "xlsx"),
        inscripciones,
        exp.COLUMNAS_PLANILLA,
        f"planilla-asistencia-{comision.curso.nombre}-{comision.fecha_inicio}",
        titulo="Asistencia",
        encabezado_extra=exp.encabezado_comision(comision)
        + ["", "PLANILLA DE ASISTENCIA — firmar al ingresar"],
    )


@requiere_panel
def exportar_inscripciones(request):
    """Exporta el listado con los filtros que el usuario dejó aplicados."""
    consulta, filtros = _filtrar_inscripciones(request)
    descripcion = [f"{clave}: {valor}" for clave, valor in filtros.items() if valor]
    return exp.exportar(
        request.GET.get("formato", "xlsx"),
        consulta,
        exp.COLUMNAS_INSCRIPCIONES,
        "inscripciones",
        titulo="Inscripciones",
        encabezado_extra=["Listado de inscripciones"]
        + ([f"Filtros — {', '.join(descripcion)}"] if descripcion else []),
    )


@requiere_panel
def exportar_participantes(request):
    """Padrón de personas (una fila por persona, no por inscripción)."""
    import csv

    from django.http import HttpResponse

    consulta = Participante.objects.select_related("empresa").annotate(
        total_cursos=Count("inscripciones", distinct=True)
    )
    busqueda = request.GET.get("q", "").strip()
    if busqueda:
        consulta = consulta.filter(
            Q(apellido__icontains=busqueda)
            | Q(nombre__icontains=busqueda)
            | Q(dni__icontains=busqueda)
        )
    empresa = request.GET.get("empresa", "")
    if empresa:
        consulta = consulta.filter(empresa_id=empresa)

    columnas = [
        "Apellido", "Nombre", "DNI", "Fecha de nacimiento", "Teléfono", "Correo",
        "Localidad", "Empresa", "Cat. licencia", "Cursos realizados", "Observaciones",
    ]
    filas = [
        [
            p.apellido, p.nombre, p.dni,
            p.fecha_nacimiento.strftime("%d/%m/%Y") if p.fecha_nacimiento else "",
            p.telefono, p.email, p.localidad,
            p.empresa.razon_social if p.empresa else "",
            p.categoria_licencia, p.total_cursos, p.observaciones,
        ]
        for p in consulta
    ]

    if request.GET.get("formato") == "csv":
        # utf-8 (no utf-8-sig): el BOM se escribe una sola vez, más abajo.
        respuesta = HttpResponse(content_type="text/csv; charset=utf-8")
        respuesta["Content-Disposition"] = (
            f'attachment; filename="padron-participantes-{date.today():%Y-%m-%d}.csv"'
        )
        respuesta.write("﻿")
        escritor = csv.writer(respuesta, delimiter=";")
        escritor.writerow(columnas)
        escritor.writerows(filas)
        return respuesta

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from django.http import HttpResponse as Respuesta

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Padrón"
    hoja.append(columnas)
    for indice, columna in enumerate(columnas, start=1):
        celda = hoja.cell(row=1, column=indice)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="0F2744")
        celda.alignment = Alignment(horizontal="center")
        hoja.column_dimensions[get_column_letter(indice)].width = max(len(columna) + 4, 14)
    for fila in filas:
        hoja.append(fila)
    hoja.freeze_panes = "A2"

    respuesta = Respuesta(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    respuesta["Content-Disposition"] = (
        f'attachment; filename="padron-participantes-{date.today():%Y-%m-%d}.xlsx"'
    )
    libro.save(respuesta)
    return respuesta


# --- Informes --------------------------------------------------------------

@requiere_panel
def informes(request):
    """Resumen por período: qué se dictó, cuánta gente pasó y qué se cobró."""
    hoy = date.today()
    desde = request.GET.get("desde") or (hoy - timedelta(days=90)).isoformat()
    hasta = request.GET.get("hasta") or hoy.isoformat()

    comisiones = Comision.objects.filter(
        fecha_inicio__gte=desde, fecha_inicio__lte=hasta
    ).select_related("curso", "curso__categoria")
    inscripciones = Inscripcion.objects.filter(
        comision__fecha_inicio__gte=desde, comision__fecha_inicio__lte=hasta
    ).exclude(estado=Inscripcion.Estado.CANCELADA)

    por_curso = (
        Curso.objects.filter(
            comisiones__fecha_inicio__gte=desde, comisiones__fecha_inicio__lte=hasta
        )
        .annotate(
            dictados=Count("comisiones", distinct=True),
            personas=Count(
                "comisiones__inscripciones",
                filter=~Q(comisiones__inscripciones__estado=Inscripcion.Estado.CANCELADA),
            ),
            cobrado=Sum(
                "comisiones__inscripciones__monto",
                filter=Q(comisiones__inscripciones__pago=Inscripcion.Pago.PAGADO),
            ),
        )
        .order_by("-personas")
    )
    tope = max([c.personas for c in por_curso] or [1]) or 1
    for curso in por_curso:
        curso.barra = round(curso.personas * 100 / tope)

    por_empresa = (
        Empresa.objects.filter(
            participantes__inscripciones__comision__fecha_inicio__gte=desde,
            participantes__inscripciones__comision__fecha_inicio__lte=hasta,
        )
        .annotate(personas=Count("participantes__inscripciones", distinct=True))
        .order_by("-personas")[:12]
    )

    return render(request, "panel/informes.html", {
        "seccion": "informes",
        "desde": desde,
        "hasta": hasta,
        "totales": {
            "comisiones": comisiones.count(),
            "personas": inscripciones.count(),
            "asistieron": inscripciones.filter(estado=Inscripcion.Estado.ASISTIO).count(),
            "cobrado": inscripciones.filter(pago=Inscripcion.Pago.PAGADO).aggregate(
                total=Sum("monto")
            )["total"] or 0,
            "pendiente_cobro": inscripciones.filter(pago=Inscripcion.Pago.PENDIENTE).count(),
        },
        "por_curso": por_curso,
        "por_empresa": por_empresa,
    })


# --- Administración --------------------------------------------------------

@SOLO_DIRECCION
def usuarios_lista(request):
    usuarios = Usuario.objects.prefetch_related("groups").order_by("first_name", "username")
    por_rol = {rol: [] for rol in JERARQUIA}
    sin_rol = []
    for usuario in usuarios:
        rol = usuario.rol
        (por_rol[rol] if rol else sin_rol).append(usuario)
    return render(request, "panel/usuarios.html", {
        "seccion": "usuarios",
        "grupos": [
            {"rol": rol, "descripcion": DESCRIPCION_ROLES[rol], "usuarios": por_rol[rol]}
            for rol in JERARQUIA
        ],
        "sin_rol": sin_rol,
        "total": usuarios.count(),
    })


@SOLO_DIRECCION
def usuario_editar(request, pk=None):
    usuario = get_object_or_404(Usuario, pk=pk) if pk else None
    if request.method == "POST":
        formulario = UsuarioForm(request.POST, instance=usuario)
        if formulario.is_valid():
            guardado = formulario.save()
            messages.success(
                request,
                f"Usuario «{guardado.username}» guardado con el rol {guardado.rol_nombre}.",
            )
            return redirect("panel:usuarios")
        messages.error(request, "Revisá los campos marcados en rojo.")
    else:
        formulario = UsuarioForm(instance=usuario)
    return render(request, "panel/usuario_form.html", {
        "seccion": "usuarios",
        "form": formulario,
        "usuario_editado": usuario,
        "descripciones": [(rol.label, DESCRIPCION_ROLES[rol]) for rol in JERARQUIA],
    })


@SOLO_DIRECCION
def configuracion(request):
    sitio = ConfiguracionSitio.vigente()
    if request.method == "POST":
        formulario = ConfiguracionForm(request.POST, instance=sitio)
        if formulario.is_valid():
            formulario.save()
            messages.success(request, "Configuración del sitio actualizada.")
            return redirect("panel:configuracion")
    else:
        formulario = ConfiguracionForm(instance=sitio)
    return render(request, "panel/configuracion.html", {
        "seccion": "configuracion",
        "form": formulario,
    })


@requiere_panel
def ayuda(request):
    """Guía de uso dentro del panel, para no depender de un manual aparte."""
    return render(request, "panel/ayuda.html", {
        "seccion": "ayuda",
        "roles": [(rol.label, DESCRIPCION_ROLES[rol]) for rol in JERARQUIA],
    })
