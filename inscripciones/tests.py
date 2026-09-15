"""Pruebas de inscripciones, exportaciones y permisos de las vistas."""
import io
from datetime import date, timedelta

from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from cuentas.models import Rol, Usuario
from cuentas.permisos import sincronizar_grupos
from cursos.models import Categoria, Comision, ConfiguracionSitio, Curso
from inscripciones.models import Empresa, Inscripcion, Participante


class BaseDatos(TestCase):
    """Un curso, una comisión y algunos usuarios de cada rol."""

    @classmethod
    def setUpTestData(cls):
        sincronizar_grupos()
        cls.categoria = Categoria.objects.create(nombre="Cargas Generales")
        cls.curso = Curso.objects.create(
            categoria=cls.categoria, nombre="Curso primera vez", precio=405000,
            requisitos="Psicofísico vigente.",
        )
        cls.comision = Comision.objects.create(
            curso=cls.curso, fecha_inicio=date.today() + timedelta(days=15),
            cupo=2, estado=Comision.Estado.PUBLICADA,
        )
        cls.usuarios = {}
        for rol in Rol:
            usuario = Usuario.objects.create_user(
                username=rol.value.lower()[:10], password="clave-larga-123",
                first_name=rol.label,
            )
            usuario.asignar_rol(rol)
            cls.usuarios[rol] = usuario

    def entrar(self, rol):
        self.client.force_login(self.usuarios[rol])


class InscripcionModeloTests(BaseDatos):
    def test_un_participante_no_puede_repetirse_en_la_misma_comision(self):
        from django.db import IntegrityError, transaction

        participante = Participante.objects.create(dni="30000001", apellido="A", nombre="B")
        Inscripcion.objects.create(comision=self.comision, participante=participante)
        with self.assertRaises(IntegrityError):
            with transaction.atomic():
                Inscripcion.objects.create(comision=self.comision, participante=participante)

    def test_pendientes_enumera_lo_que_falta(self):
        participante = Participante.objects.create(dni="30000002", apellido="A", nombre="B")
        inscripcion = Inscripcion.objects.create(comision=self.comision, participante=participante)
        self.assertEqual(inscripcion.pendientes, ["pago", "psicofísico", "documentación"])

        inscripcion.pago = Inscripcion.Pago.PAGADO
        inscripcion.psicofisico_vigente = True
        inscripcion.documentacion_completa = True
        self.assertEqual(inscripcion.pendientes, [])
        inscripcion.estado = Inscripcion.Estado.CONFIRMADA
        self.assertTrue(inscripcion.lista_para_cursar)


class InscribirVistaTests(BaseDatos):
    def url(self):
        return reverse("panel:inscribir", kwargs={"pk": self.comision.pk})

    def test_recepcion_puede_inscribir_a_alguien_nuevo(self):
        self.entrar(Rol.RECEPCION)
        respuesta = self.client.post(self.url(), {
            "dni": "31000001", "apellido": "Gómez", "nombre": "Juan",
            "telefono": "343 4000000", "email": "", "empresa": "",
            "estado": Inscripcion.Estado.CONFIRMADA, "pago": Inscripcion.Pago.PENDIENTE,
        })
        self.assertEqual(respuesta.status_code, 302)
        self.assertTrue(Participante.objects.filter(dni="31000001").exists())
        self.assertEqual(self.comision.cantidad_inscriptos, 1)

    def test_reutiliza_a_la_persona_si_el_dni_ya_existe(self):
        Participante.objects.create(dni="31000002", apellido="Sosa", nombre="Ana")
        self.entrar(Rol.RECEPCION)
        self.client.post(self.url(), {
            "dni": "31000002", "apellido": "", "nombre": "", "telefono": "", "email": "",
            "empresa": "", "estado": Inscripcion.Estado.CONFIRMADA,
            "pago": Inscripcion.Pago.PENDIENTE,
        })
        self.assertEqual(Participante.objects.filter(dni="31000002").count(), 1)
        self.assertEqual(self.comision.cantidad_inscriptos, 1)

    def test_un_dni_nuevo_sin_nombre_es_rechazado(self):
        self.entrar(Rol.RECEPCION)
        self.client.post(self.url(), {
            "dni": "31000003", "apellido": "", "nombre": "", "telefono": "", "email": "",
            "empresa": "", "estado": Inscripcion.Estado.CONFIRMADA,
            "pago": Inscripcion.Pago.PENDIENTE,
        })
        self.assertFalse(Participante.objects.filter(dni="31000003").exists())

    def test_no_se_puede_pasar_del_cupo(self):
        self.entrar(Rol.RECEPCION)
        for dni in ("31000010", "31000011"):
            self.client.post(self.url(), {
                "dni": dni, "apellido": "X", "nombre": "Y", "telefono": "", "email": "",
                "empresa": "", "estado": Inscripcion.Estado.CONFIRMADA,
                "pago": Inscripcion.Pago.PENDIENTE,
            })
        self.assertTrue(self.comision.completa)
        self.client.post(self.url(), {
            "dni": "31000012", "apellido": "Z", "nombre": "W", "telefono": "", "email": "",
            "empresa": "", "estado": Inscripcion.Estado.CONFIRMADA,
            "pago": Inscripcion.Pago.PENDIENTE,
        })
        self.assertEqual(self.comision.cantidad_inscriptos, 2)

    def test_un_instructor_no_puede_inscribir(self):
        self.entrar(Rol.INSTRUCTOR)
        self.client.post(self.url(), {
            "dni": "31000020", "apellido": "X", "nombre": "Y", "telefono": "", "email": "",
            "empresa": "", "estado": Inscripcion.Estado.CONFIRMADA,
            "pago": Inscripcion.Pago.PENDIENTE,
        })
        self.assertEqual(self.comision.cantidad_inscriptos, 0)


class PermisosTests(BaseDatos):
    def test_el_panel_exige_ingresar(self):
        respuesta = self.client.get(reverse("panel:inicio"))
        self.assertEqual(respuesta.status_code, 302)
        self.assertIn("ingresar", respuesta["Location"])

    def test_solo_direccion_administra_usuarios(self):
        self.entrar(Rol.DIRECCION)
        self.assertEqual(self.client.get(reverse("panel:usuarios")).status_code, 200)
        for rol in (Rol.ADMINISTRACION, Rol.COORDINACION, Rol.RECEPCION, Rol.INSTRUCTOR):
            self.entrar(rol)
            self.assertEqual(
                self.client.get(reverse("panel:usuarios")).status_code, 302,
                f"{rol.value} no debería entrar a Usuarios",
            )

    def test_coordinacion_no_edita_el_catalogo(self):
        url = reverse("panel:curso_editar", kwargs={"pk": self.curso.pk})
        self.entrar(Rol.ADMINISTRACION)
        self.assertEqual(self.client.get(url).status_code, 200)
        self.entrar(Rol.COORDINACION)
        self.assertEqual(self.client.get(url).status_code, 302)

    def test_recepcion_no_crea_comisiones(self):
        self.entrar(Rol.COORDINACION)
        self.assertEqual(self.client.get(reverse("panel:comision_nueva")).status_code, 200)
        self.entrar(Rol.RECEPCION)
        self.assertEqual(self.client.get(reverse("panel:comision_nueva")).status_code, 302)

    def test_el_instructor_solo_ve_sus_comisiones(self):
        otra = Comision.objects.create(
            curso=self.curso, fecha_inicio=date.today() + timedelta(days=30), cupo=10
        )
        self.comision.instructor = self.usuarios[Rol.INSTRUCTOR]
        self.comision.save()
        self.entrar(Rol.INSTRUCTOR)

        propia = reverse("panel:comision_detalle", kwargs={"pk": self.comision.pk})
        ajena = reverse("panel:comision_detalle", kwargs={"pk": otra.pk})
        self.assertEqual(self.client.get(propia).status_code, 200)
        self.assertEqual(self.client.get(ajena).status_code, 404)


class ExportacionTests(BaseDatos):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        empresa = Empresa.objects.create(razon_social="Transporte Litoral SRL")
        for indice, apellido in enumerate(("Gómez", "Álvarez")):
            participante = Participante.objects.create(
                dni=f"3200000{indice}", apellido=apellido, nombre="Juan",
                telefono="343 4000000", empresa=empresa,
            )
            Inscripcion.objects.create(
                comision=cls.comision, participante=participante,
                estado=Inscripcion.Estado.CONFIRMADA, pago=Inscripcion.Pago.PAGADO,
                monto=405000,
            )

    def test_excel_trae_encabezado_y_una_fila_por_participante(self):
        self.entrar(Rol.RECEPCION)
        respuesta = self.client.get(
            reverse("panel:exportar_comision", kwargs={"pk": self.comision.pk}),
            {"formato": "xlsx"},
        )
        self.assertEqual(respuesta.status_code, 200)
        self.assertIn("spreadsheetml", respuesta["Content-Type"])
        self.assertIn("attachment; filename=", respuesta["Content-Disposition"])

        hoja = load_workbook(io.BytesIO(respuesta.content)).active
        textos = [str(c.value) for c in hoja["A"] if c.value]
        self.assertTrue(any("Curso: Curso primera vez" in t for t in textos))
        self.assertTrue(any("Cupo: 2 de 2" in t for t in textos))
        # Encabezado de la tabla + las dos personas, ordenadas por apellido.
        apellidos = [
            hoja.cell(row=f, column=2).value for f in range(1, hoja.max_row + 1)
        ]
        self.assertIn("Álvarez", apellidos)
        self.assertIn("Gómez", apellidos)

    def test_csv_lleva_un_unico_bom_y_separador_punto_y_coma(self):
        self.entrar(Rol.RECEPCION)
        respuesta = self.client.get(
            reverse("panel:exportar_comision", kwargs={"pk": self.comision.pk}),
            {"formato": "csv"},
        )
        self.assertEqual(respuesta.status_code, 200)
        # Un solo BOM: si no, Excel muestra basura al principio de cada línea.
        self.assertEqual(respuesta.content.count(b"\xef\xbb\xbf"), 1)
        texto = respuesta.content.decode("utf-8-sig")
        self.assertIn("Apellido;Nombre;DNI", texto)
        self.assertIn("Gómez", texto)

    def test_la_planilla_de_firmas_deja_las_columnas_en_blanco(self):
        self.entrar(Rol.RECEPCION)
        respuesta = self.client.get(
            reverse("panel:planilla", kwargs={"pk": self.comision.pk}), {"formato": "xlsx"}
        )
        hoja = load_workbook(io.BytesIO(respuesta.content)).active
        encabezados = [
            c.value for f in range(1, hoja.max_row + 1)
            for c in [hoja.cell(row=f, column=5)] if c.value
        ]
        self.assertIn("Firma", encabezados)

    def test_las_canceladas_quedan_fuera_de_la_lista(self):
        inscripcion = self.comision.inscripciones.first()
        inscripcion.estado = Inscripcion.Estado.CANCELADA
        inscripcion.save()
        self.entrar(Rol.RECEPCION)
        respuesta = self.client.get(
            reverse("panel:exportar_comision", kwargs={"pk": self.comision.pk}),
            {"formato": "csv"},
        )
        texto = respuesta.content.decode("utf-8-sig")
        self.assertNotIn(inscripcion.participante.apellido, texto)

    def test_la_exportacion_respeta_los_filtros_de_pantalla(self):
        self.entrar(Rol.ADMINISTRACION)
        respuesta = self.client.get(
            reverse("panel:exportar_inscripciones"), {"formato": "csv", "pago": "pendiente"}
        )
        texto = respuesta.content.decode("utf-8-sig")
        # Las dos inscripciones están pagadas, así que no debe salir ninguna.
        self.assertNotIn("Gómez", texto)


class WebPublicaTests(BaseDatos):
    def test_el_inicio_muestra_los_cursos_activos(self):
        respuesta = self.client.get(reverse("web:inicio"))
        self.assertEqual(respuesta.status_code, 200)
        self.assertContains(respuesta, "Curso primera vez")
        self.assertContains(respuesta, "$405.000")

    def test_un_curso_desactivado_no_aparece(self):
        self.curso.activo = False
        self.curso.save()
        respuesta = self.client.get(reverse("web:inicio"))
        self.assertNotContains(respuesta, "Curso primera vez")

    def test_la_ficha_del_curso_lista_las_fechas_publicadas(self):
        respuesta = self.client.get(reverse("web:curso", kwargs={"slug": self.curso.slug}))
        self.assertEqual(respuesta.status_code, 200)
        self.assertContains(respuesta, self.comision.fecha_texto)

    def test_la_preinscripcion_crea_persona_e_inscripcion(self):
        respuesta = self.client.post(
            reverse("web:preinscripcion", kwargs={"slug": self.curso.slug}),
            {
                "comision": self.comision.pk, "dni": "33000001", "apellido": "Ledesma",
                "nombre": "Carlos", "telefono": "343 4111111", "email": "c@ejemplo.com",
                "localidad": "Paraná", "empresa": "Fletes del Sur",
            },
        )
        self.assertEqual(respuesta.status_code, 302)
        participante = Participante.objects.get(dni="33000001")
        inscripcion = Inscripcion.objects.get(participante=participante)
        self.assertEqual(inscripcion.estado, Inscripcion.Estado.PREINSCRIPTO)
        self.assertEqual(inscripcion.origen, "web")
        self.assertEqual(participante.empresa.razon_social, "Fletes del Sur")

    def test_no_se_puede_reservar_en_una_comision_completa(self):
        for indice in range(2):
            participante = Participante.objects.create(
                dni=f"3400000{indice}", apellido="X", nombre="Y"
            )
            Inscripcion.objects.create(comision=self.comision, participante=participante)
        self.client.post(
            reverse("web:preinscripcion", kwargs={"slug": self.curso.slug}),
            {
                "comision": self.comision.pk, "dni": "33000009", "apellido": "Tarde",
                "nombre": "Llego", "telefono": "343 4111111", "email": "",
                "localidad": "", "empresa": "",
            },
        )
        self.assertFalse(Participante.objects.filter(dni="33000009").exists())

    def test_la_preinscripcion_se_puede_apagar_desde_la_configuracion(self):
        sitio = ConfiguracionSitio.vigente()
        sitio.inscripcion_online = False
        sitio.save()
        respuesta = self.client.get(
            reverse("web:preinscripcion", kwargs={"slug": self.curso.slug})
        )
        self.assertEqual(respuesta.status_code, 302)


class AsistenciaTests(BaseDatos):
    def test_tomar_asistencia_marca_presentes_y_ausentes(self):
        inscripciones = []
        for indice in range(2):
            participante = Participante.objects.create(
                dni=f"3500000{indice}", apellido=f"A{indice}", nombre="B"
            )
            inscripciones.append(
                Inscripcion.objects.create(
                    comision=self.comision, participante=participante,
                    estado=Inscripcion.Estado.CONFIRMADA,
                )
            )
        self.entrar(Rol.DIRECCION)
        self.client.post(
            reverse("panel:asistencia", kwargs={"pk": self.comision.pk}),
            {"presente": [inscripciones[0].pk]},
        )
        inscripciones[0].refresh_from_db()
        inscripciones[1].refresh_from_db()
        self.comision.refresh_from_db()
        self.assertEqual(inscripciones[0].estado, Inscripcion.Estado.ASISTIO)
        self.assertEqual(inscripciones[1].estado, Inscripcion.Estado.AUSENTE)
        self.assertEqual(self.comision.estado, Comision.Estado.FINALIZADA)
